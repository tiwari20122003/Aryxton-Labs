import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def _header_style(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF8B7FD3")
        cell.alignment = Alignment(horizontal="left", vertical="center")


def generate_excel_report(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"

    domain = payload.get("domain", {})
    cyber = payload.get("cyber", {})
    seo = payload.get("seo", {})

    ws.append(["Website Analyser Report", ""])
    ws.append(["Generated", datetime.utcnow().isoformat()])
    ws.append([])
    ws.append(["Metric", "Value"])
    overview_rows = [
        ("Domain", domain.get("domain", "N/A")),
        ("Registrar", domain.get("registrar", "N/A")),
        ("Creation Date", domain.get("creation_date", "N/A")),
        ("Expiration Date", domain.get("expiration_date", "N/A")),
        ("Name Servers", domain.get("name_servers", "N/A")),
        ("MX Records", domain.get("mx_records", "N/A")),
        ("SSL Status", domain.get("ssl_status", "N/A")),
        ("robots.txt", domain.get("robots_txt", "N/A")),
        ("sitemap.xml", domain.get("sitemap_xml", "N/A")),
        ("Security Grade", cyber.get("grade", "N/A")),
        ("Security Score", cyber.get("global_score", "N/A")),
        ("Pages Crawled", len(seo.get("pages", []))),
        ("SEO Issues", len(seo.get("issues", []))),
    ]
    for k, v in overview_rows:
        ws.append([k, v])
    _header_style(ws, row=4)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60

    ws2 = wb.create_sheet("Pages")
    ws2.append(["URL", "Title", "Meta Description", "H1", "Status"])
    for p in seo.get("pages", []):
        ws2.append([p.get("url"), p.get("title"), p.get("meta_description"), p.get("h1"), p.get("status")])
    _header_style(ws2)
    for col in ["A", "B", "C", "D", "E"]:
        ws2.column_dimensions[col].width = 32
    ws3 = wb.create_sheet("Issues")
    ws3.append(["URL", "Issue", "Severity", "Category", "Description", "Fix"])
    for i in seo.get("issues", []):
        ws3.append([i.get("url"), i.get("issue"), i.get("severity"), i.get("category"), i.get("description"), i.get("fix")])
    _header_style(ws3)
    for col in ["A", "B", "C", "D", "E", "F"]:
        ws3.column_dimensions[col].width = 30
    ws4 = wb.create_sheet("Performance Audit")
    ws4.append(["URL", "Load Time (s)", "Page Size (KB)", "Missing Alt", "Canonical", "Social Meta", "Structured Data",
                "Internal Links", "External Links", "Title Len", "Meta Len", "LCP (s)", "Page Speed Score", "LCP Target"])
    for a in seo.get("audits", []):
        ws4.append([a.get("url"), a.get("load_time_sec"), a.get("page_size_kb"), a.get("missing_alt_images"),
                    a.get("canonical_tag"), a.get("social_meta_og"), a.get("structured_data"),
                    a.get("internal_links"), a.get("external_links"), a.get("title_length"),
                    a.get("meta_length"), a.get("lcp_sec"), a.get("page_speed_score"), a.get("lcp_target")])
    _header_style(ws4)

    ws5 = wb.create_sheet("Security")
    ws5.append(["Metric", "Value"])
    sec_rows = [
        ("Grade", cyber.get("grade")),
        ("Score", cyber.get("global_score")),
        ("Rating", cyber.get("rating")),
        ("Headers Score", cyber.get("headers_score")),
        ("Phishing Score", cyber.get("phishing_score")),
        ("Malware Score", cyber.get("malware_score")),
        ("SSL Valid", cyber.get("ssl_info", {}).get("valid")),
        ("SSL Issuer", cyber.get("ssl_info", {}).get("issuer_org")),
        ("SSL Expiry", cyber.get("ssl_info", {}).get("expiry_date")),
        ("TLS Version", cyber.get("ssl_info", {}).get("tls_version")),
    ]
    for k, v in sec_rows:
        ws5.append([k, v])
    _header_style(ws5)
    ws5.append([])
    ws5.append(["Header", "Status", "Severity", "Description"])
    _header_style(ws5, row=len(sec_rows) + 3)
    for f in cyber.get("header_findings", []):
        ws5.append([f.get("header"), f.get("status"), f.get("severity"), f.get("desc")])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
